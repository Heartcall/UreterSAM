from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "runs" / "UreterSAM_当前技术方案.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, bold=False, size=10.5):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return run


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 11.5, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_text(p, text)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, bold=True, size=9.5).font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_text(p, str(value), size=9.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    add_text(title, "UreterSAM 当前整体技术方案", bold=True, size=20)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    add_text(subtitle, "输尿管镜视频中的结石分割 深度估计 尺寸计算与激光准星预测", size=10.5)

    add_body(doc, "本文说明当前 UreterSAM 代码的完整技术流程，包括结石和光纤分割、单目深度估计、光纤直径标定、最近结石选择、结石尺寸与体积估计，以及激光预计照射点的计算方法。当前系统输出的是基于视觉和相对深度的工程估计结果，准星和尺寸仍需要通过真实激光落点及标准尺寸样本进一步标定。")

    add_heading(doc, "一 系统整体流程", 1)
    add_body(doc, "系统从输尿管镜视频或单帧图像开始，先使用双 U-Net 分割模型得到结石和光纤掩膜，再结合相机内参和 Depth Anything V2 深度结果建立几何关系。程序随后选择光纤激光方向上最近的结石，使用光纤末端直径进行尺度换算，并输出带有分割、中心点、尺寸标签和绿色准星的视频或图像。")
    add_table(doc, ["阶段", "主要处理", "输出"], [
        ("输入", "读取视频帧或 JPG/PNG 图像并统一预处理", "标准化图像"),
        ("分割", "双 U-Net 融合预测背景、结石和光纤", "三类分割掩膜"),
        ("几何处理", "连通域分析、相机参数缩放和畸变处理", "结石与光纤几何信息"),
        ("深度估计", "Depth Anything V2 Small 输出单目相对深度", "深度图"),
        ("测量", "使用光纤直径标定并估计最近结石尺寸", "毫米尺寸和体积"),
        ("准星", "根据光纤方向和深度估计激光预计落点", "绿色准星坐标"),
    ], widths=[2.2, 8.0, 4.0])

    add_heading(doc, "二 分割模型与训练方法", 1)
    add_body(doc, "分割网络采用 ResNet34 编码器和 U-Net 解码器，输出背景、结石和光纤三个类别。网络定义在 ureter_semis/model.py，底层 U-Net 结构位于 third_party/Endo-SemiS/src/models/unet.py。")
    add_bullet(doc, "解码器中加入 Dropout2d，用于提高泛化能力，并支持通过 Monte Carlo Dropout 估计预测不确定性。")
    add_bullet(doc, "训练阶段同时使用两个结构相同的 U-Net。两个模型根据预测置信度和熵生成伪标签，并对无标签图像进行交叉监督。")
    add_bullet(doc, "有标签数据使用带类别权重的 Cross Entropy 和多分类 Dice Loss；无标签数据使用高置信度伪标签和一致性损失。")
    add_bullet(doc, "推理阶段对两个模型的 logits 进行平均，再取最大类别作为最终分割结果。")
    add_body(doc, "当前输出类别为：0 表示背景，1 表示结石，2 表示光纤。结果图中红色区域表示结石，青色区域表示光纤。")

    add_heading(doc, "三 使用的数据集", 1)
    add_body(doc, "训练数据来自 Roboflow Universe 用户 urology 发布的多个公开实例分割数据集。项目将四个数据集统一转换为背景、结石和光纤三类，并合并为 data/kidney_stone_fiber_all。部分原始数据还包含 tissue 或 laser 类别，当前转换流程没有将这些类别作为模型输出类别。")
    add_table(doc, ["数据集", "原始类别特点", "当前处理"], [
        ("Kidney Stone Fiber", "stone、fiber 及其他类别", "保留 stone 和 fiber"),
        ("kidney stone fiber tissue", "stone、fiber、tissue", "忽略 tissue"),
        ("Kidney stones fiber", "stone、fiber、tissue、laser", "保留 stone 和 fiber"),
        ("kidney stone fiber tissue 060125", "stone、fiber、tissue", "忽略 tissue"),
    ], widths=[5.0, 6.2, 3.0])

    add_heading(doc, "四 深度估计", 1)
    add_body(doc, "深度估计代码位于 ureter_semis/depth_anything_v2_infer.py，默认调用 Hugging Face 上的 Depth Anything V2 Small 模型。模型对每一帧图像生成单通道深度结果，用于比较光纤和结石之间的相对空间关系。")
    add_bullet(doc, "当前深度结果属于单目相对深度或逆深度形式，不是直接以毫米表示的绝对距离。")
    add_bullet(doc, "深度估计主要用于相对距离修正、光纤方向上的空间点搜索和准星预测。")
    add_bullet(doc, "结石尺寸的毫米单位主要由已知光纤末端直径提供，深度图用于辅助修正不同深度位置产生的尺度变化。")

    add_heading(doc, "五 光纤尺寸标定", 1)
    add_body(doc, "系统默认已知光纤末端直径为 0.4 mm。程序从光纤分割掩膜中估计光纤的像素宽度，并计算像素到毫米的比例：")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "毫米每像素 = 0.4 ÷ 光纤分割宽度像素值", bold=True, size=11)
    add_body(doc, "在视频测试中还可以提供光纤可见伸出长度 4 mm，使系统同时参考光纤直径和光纤长度进行尺度估计。单帧图像测量默认主要按照末端直径 0.4 mm 进行标定。")

    add_heading(doc, "六 最近结石的选择与尺寸估计", 1)
    add_body(doc, "系统会分割图像中的所有结石，但当前尺寸估计只针对光纤激光方向上最近的结石。程序先通过光纤连通域估计光纤轴线和传播方向，然后检查该方向上的结石连通域，选择最接近光纤的结石作为测量对象。其他结石仍然显示分割结果，但不参与当前尺寸标签。")
    add_body(doc, "结石像素尺寸经过光纤尺度换算后，得到长度和宽度。体积目前采用椭球近似，公式为：")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "体积 = π ÷ 6 × 长度 × 宽度 × 宽度", bold=True, size=11)
    add_body(doc, "由于单幅内窥镜图像无法直接观察结石的完整第三空间方向，程序暂时用结石宽度近似第三方向。因此体积是估计值，需要使用带真实尺寸的结石样本进行验证。")

    add_heading(doc, "七 激光准星计算", 1)
    add_body(doc, "绿色准星表示预计的激光照射点。准星的计算不再直接取结石中心，也不强制限制在结石分割区域内部。结石分割只用于选择需要测量的最近结石，准星则根据光纤方向和深度信息独立计算。")
    add_table(doc, ["步骤", "处理逻辑"], [
        ("1", "从光纤分割掩膜中提取最大光纤连通区域。"),
        ("2", "使用 PCA 估计光纤中心轴线。"),
        ("3", "根据光纤在图像中的位置判断近端和末端，建立从近端指向末端的方向。"),
        ("4", "使用相机内参将光纤近端和末端结合深度结果反投影到近似三维坐标。"),
        ("5", "沿光纤末端方向在图像中向前采样，并读取对应深度。"),
        ("6", "在三维空间中寻找距离光纤激光射线最近的采样点。"),
        ("7", "将该空间点重新投影到图像平面，作为绿色准星位置。"),
    ], widths=[1.5, 12.7])
    add_body(doc, "如果深度图不可用，程序会退化为沿光纤方向延伸到图像边界的位置。准星可能落在结石上，也可能落在结石外部，这取决于光纤方向、深度估计和当前视野中的空间关系。")

    add_heading(doc, "八 相机参数与几何校正", 1)
    add_body(doc, "程序读取 camera_parameters.txt 中提供的相机内参和畸变参数。参考分辨率为 1422×1066，实际测试视频通常为 400×400，因此当前代码会根据实际图像尺寸对焦距和主点进行缩放，并使用畸变参数进行几何校正。")
    add_body(doc, "如果 400×400 图像是由原始画面非等比例缩放得到，当前缩放方法可以作为近似；如果 400×400 是裁剪、居中裁剪或经过其他视口变换，则还需要根据真实图像形成过程重新计算相机内参。")

    add_heading(doc, "九 输出结果", 1)
    add_table(doc, ["输出", "内容"], [
        ("带标注视频或图像", "红色结石分割、青色光纤分割、黄色结石中心点、绿色激光预计照射点和尺寸文字。"),
        ("JSONL 测量文件", "每帧或每张图像的结石尺寸、体积、深度、光纤方向和准星坐标。"),
        ("模型权重", "runs/endosemis_256_all_5ep/best.pt。"),
    ], widths=[4.0, 10.2])

    add_heading(doc, "十 当前技术边界", 1)
    add_bullet(doc, "Depth Anything V2 提供相对深度，当前没有建立严格的毫米级深度标定。")
    add_bullet(doc, "准星是根据二维光纤分割、单目深度和相机内参推断的视觉结果，尚未使用真实激光落点进行系统标定。")
    add_bullet(doc, "结石体积采用椭球近似，第三空间方向使用宽度近似。")
    add_bullet(doc, "如果需要临床或工程精度，下一步应采集带真实结石尺寸、真实激光落点和准确成像分辨率的标定数据。")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "UreterSAM 技术方案", size=9)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
